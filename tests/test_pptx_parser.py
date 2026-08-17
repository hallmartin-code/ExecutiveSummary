"""PowerPoint deck ingestion."""

from __future__ import annotations

from pathlib import Path

import pytest

from src.ingestion import load_deck
from src.ingestion.pdf_parser import DeckParseError
from src.ingestion.pptx_parser import parse_pptx_deck


def test_parses_every_slide(sample_pptx_deck: Path) -> None:
    deck = parse_pptx_deck(sample_pptx_deck)
    assert deck.source_format == "pptx"
    assert deck.slide_count == 9
    assert [s.slide for s in deck.slides] == list(range(1, 10))


def test_titles_and_body_are_separated(sample_pptx_deck: Path) -> None:
    deck = parse_pptx_deck(sample_pptx_deck)
    problem = deck.slides[1]
    assert problem.title == "Problem"
    assert "14 hours per week" in problem.content


def test_speaker_notes_are_extracted(sample_pptx_deck: Path) -> None:
    deck = parse_pptx_deck(sample_pptx_deck)
    assert deck.slides[0].speaker_notes is not None
    assert "Speaker notes" in deck.slides[0].speaker_notes


def test_tables_are_extracted(sample_pptx_deck: Path) -> None:
    deck = parse_pptx_deck(sample_pptx_deck)
    market = deck.slides[3]
    assert market.tables, "expected the market slide's table to be captured"
    flattened = market.tables[0].to_text()
    assert "TAM" in flattened and "$12.5B" in flattened


def test_extra_shapes_are_captured(sample_pptx_deck: Path) -> None:
    """Text in additional shapes must not be lost."""
    deck = parse_pptx_deck(sample_pptx_deck)
    traction = deck.slides[5]
    assert "120 qualified shippers" in traction.all_text()


def test_metrics_carry_slide_numbers(sample_pptx_deck: Path) -> None:
    deck = parse_pptx_deck(sample_pptx_deck)
    for slide in deck.slides:
        assert all(m.slide == slide.slide for m in slide.metrics)


def test_loader_dispatches_pptx(sample_pptx_deck: Path) -> None:
    assert load_deck(sample_pptx_deck).source_format == "pptx"


def test_pptx_and_pdf_agree_on_key_facts(sample_pptx_deck: Path, sample_pdf_deck: Path) -> None:
    """The same deck in either format must yield the same headline figures."""
    pptx_text = parse_pptx_deck(sample_pptx_deck).full_text()
    pdf_text = load_deck(sample_pdf_deck).full_text()
    for token in ("$3.2M", "47 customers", "Raising $8M", "74%"):
        assert token in pptx_text, f"{token} missing from PPTX extraction"
        assert token in pdf_text, f"{token} missing from PDF extraction"


def test_malformed_pptx_raises_parse_error(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pptx"
    broken.write_bytes(b"PK\x03\x04" + b"\x00" * 200)
    with pytest.raises(DeckParseError):
        parse_pptx_deck(broken)


def test_empty_presentation_raises(tmp_path: Path) -> None:
    from pptx import Presentation

    path = tmp_path / "empty.pptx"
    Presentation().save(str(path))
    with pytest.raises(DeckParseError, match="no slides"):
        parse_pptx_deck(path)
