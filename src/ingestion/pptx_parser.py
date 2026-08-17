"""PowerPoint pitch-deck parsing via python-pptx.

Grouped shapes are walked recursively, since decks routinely bury metric
callouts inside nested groups where a flat shape iteration would miss them.
"""

from __future__ import annotations

from pathlib import Path
from typing import Iterator

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..config import MAX_DECK_PAGES
from ..utils.logging import get_logger
from .document import (
    DeckDocument,
    SlideContent,
    SlideTable,
    clean_text,
    find_metrics,
    split_bullets,
)
from .pdf_parser import DeckParseError

log = get_logger("pptx_parser")


def _walk(shapes) -> Iterator[object]:
    """Yield every shape, descending into groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            try:
                yield from _walk(shape.shapes)
            except AttributeError:  # pragma: no cover - malformed group
                continue
        else:
            yield shape


def _shape_text(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    try:
        return clean_text(shape.text_frame.text)
    except Exception:  # pragma: no cover
        return ""


def _table_rows(shape, slide_no: int) -> SlideTable | None:
    if not getattr(shape, "has_table", False):
        return None
    try:
        rows = [
            [clean_text(cell.text) for cell in row.cells]
            for row in shape.table.rows
        ]
    except Exception:  # pragma: no cover
        return None
    rows = [r for r in rows if any(c for c in r)]
    return SlideTable(rows=rows, slide=slide_no) if len(rows) >= 2 else None


def _chart_labels(shape) -> list[str]:
    """Pull category names, series names, and values out of a native chart."""
    if not getattr(shape, "has_chart", False):
        return []
    labels: list[str] = []
    try:
        chart = shape.chart
        try:
            for cat in chart.plots[0].categories:
                if cat is not None:
                    labels.append(clean_text(str(cat)))
        except (IndexError, AttributeError, ValueError):
            pass
        for series in chart.series:
            name = getattr(series, "name", None)
            if name:
                labels.append(clean_text(str(name)))
            try:
                values = [v for v in series.values if v is not None]
                if values:
                    labels.append(
                        f"{name or 'series'}: "
                        + ", ".join(f"{v:g}" for v in values[:12])
                    )
            except (AttributeError, TypeError):
                continue
    except Exception as exc:  # pragma: no cover - chart XML varies wildly
        log.debug("Chart extraction failed: %s", exc)
    return [l for l in labels if l]


def _title_of(slide, texts: list[str]) -> str | None:
    """Prefer the placeholder title; fall back to the first short text run."""
    try:
        if slide.shapes.title is not None:
            t = clean_text(slide.shapes.title.text)
            if t:
                return t
    except Exception:  # pragma: no cover
        pass
    for t in texts:
        first = t.split("\n")[0].strip()
        if 3 <= len(first) <= 90:
            return first
    return None


def parse_pptx_deck(path: str | Path) -> DeckDocument:
    """Parse a .pptx pitch deck into slide-level content."""
    p = Path(path)
    try:
        prs = Presentation(str(p))
    except Exception as exc:
        raise DeckParseError(f"Could not open PowerPoint file '{p.name}': {exc}") from exc

    slide_list = list(prs.slides)
    if not slide_list:
        raise DeckParseError(f"'{p.name}' contains no slides.")
    if len(slide_list) > MAX_DECK_PAGES:
        raise DeckParseError(
            f"'{p.name}' has {len(slide_list)} slides, above the {MAX_DECK_PAGES}-slide limit."
        )

    slides: list[SlideContent] = []
    for index, slide in enumerate(slide_list):
        slide_no = index + 1
        texts: list[str] = []
        tables: list[SlideTable] = []
        chart_labels: list[str] = []
        captions: list[str] = []
        picture_count = 0

        for shape in _walk(slide.shapes):
            text = _shape_text(shape)
            if text:
                texts.append(text)

            table = _table_rows(shape, slide_no)
            if table:
                tables.append(table)

            chart_labels.extend(_chart_labels(shape))

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                picture_count += 1
                # Alt text doubles as an image caption when the author set one.
                try:
                    alt = clean_text(shape._element._nvXxPr.cNvPr.get("descr", "") or "")
                except Exception:
                    alt = ""
                if alt:
                    captions.append(alt)

        notes = None
        try:
            if slide.has_notes_slide:
                notes = clean_text(slide.notes_slide.notes_text_frame.text) or None
        except Exception:  # pragma: no cover
            notes = None

        body = "\n".join(texts)
        title = _title_of(slide, texts)
        if title and body.startswith(title):
            body = clean_text(body[len(title):])

        combined = "\n".join([body, *chart_labels, *(t.to_text() for t in tables)])
        text_chars = len(combined)
        # Without a raster we approximate density from picture count vs. text volume.
        density = round(
            min(picture_count * 0.35, 1.2) + (1.0 - min(text_chars / 700.0, 1.0)) * 0.65,
            3,
        )

        slides.append(
            SlideContent(
                slide=slide_no,
                title=title,
                content=body,
                bullets=split_bullets(body),
                tables=tables,
                chart_labels=chart_labels[:40],
                image_captions=captions,
                speaker_notes=notes,
                metrics=find_metrics(combined, slide_no),
                image_count=picture_count,
                native_text_chars=text_chars,
                visual_density=density,
                needs_visual_pass=(text_chars < 420 and picture_count > 0) or density > 0.85,
            )
        )

    metadata: dict[str, str] = {}
    try:
        props = prs.core_properties
        for key in ("title", "author", "subject", "comments"):
            value = getattr(props, key, None)
            if value:
                metadata[key] = str(value)
    except Exception:  # pragma: no cover
        pass

    log.info("Parsed %d-slide PPTX deck: %s", len(slides), p.name)
    return DeckDocument(
        source_path=str(p),
        source_format="pptx",
        slide_count=len(slides),
        slides=slides,
        metadata=metadata,
    )
